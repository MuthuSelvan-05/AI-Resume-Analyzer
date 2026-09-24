from pathlib import Path

from docx import Document


def extract_text_from_docx(file_path: str) -> str:
    """
    Extract text from paragraphs and tables in a DOCX resume.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    document = Document(str(path))

    text_parts = []

    # Extract normal paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    # Extract text from tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                text_parts.append(" | ".join(row_text))

    return "\n".join(text_parts).strip()